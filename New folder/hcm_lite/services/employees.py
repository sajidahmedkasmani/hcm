"""M3 — Employees: registration (creates the login), directory & search, Employee 360,
self-service profile edits, exits, and CSV import (with dry run) / export."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from sqlalchemy import func, or_, select
from sqlalchemy.orm import joinedload

from ..context import Ctx, assert_can_view_employee, require_roles
from ..errors import Conflict, NotFound, ValidationFailed
from ..models import Department, Designation, Employee, EmploymentStatus, Role, User
from ..security import generate_temp_password, hash_password
from . import audit, salary
from .common import paginate, snapshot

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


@dataclass
class EmployeeInput:
    full_name: str
    email: str
    department_id: int
    join_date: date
    basic_salary: float | None = None
    designation_id: int | None = None
    employee_code: str | None = None
    gender: str | None = None
    phone: str | None = None
    address: str | None = None
    status: str = EmploymentStatus.ACTIVE
    extra: dict[str, Any] = field(default_factory=dict)


# ------------------------------------------------------------------ helpers
def _next_code(ctx: Ctx) -> str:
    n = (ctx.db.scalar(select(func.count(Employee.id))) or 0) + 1
    while True:
        code = f"EMP-{n:04d}"
        if not ctx.db.scalar(select(Employee.id).where(Employee.employee_code == code)):
            return code
        n += 1


def _validate_input(ctx: Ctx, data: EmployeeInput, *, require_salary: bool) -> None:
    if not data.full_name.strip():
        raise ValidationFailed("full_name is required", details={"field": "full_name"})
    if not EMAIL_RE.match(data.email.strip()):
        raise ValidationFailed("A valid email is required", details={"field": "email"})
    email = data.email.strip().lower()
    if ctx.db.scalar(select(User.id).where(User.email == email)) or ctx.db.scalar(
        select(Employee.id).where(Employee.email == email)
    ):
        raise Conflict(f"Email {email} is already registered")
    if ctx.db.get(Department, data.department_id) is None:
        raise ValidationFailed("Unknown department", details={"field": "department_id"})
    if data.designation_id is not None and ctx.db.get(Designation, data.designation_id) is None:
        raise ValidationFailed("Unknown designation", details={"field": "designation_id"})
    if data.employee_code and ctx.db.scalar(
        select(Employee.id).where(Employee.employee_code == data.employee_code.strip())
    ):
        raise Conflict(f"Employee code {data.employee_code} already exists")
    if data.status not in (EmploymentStatus.ACTIVE, EmploymentStatus.PROBATION):
        raise ValidationFailed("status must be 'active' or 'probation'")
    if require_salary and (data.basic_salary is None or data.basic_salary <= 0):
        raise ValidationFailed("basic_salary must be greater than 0", details={"field": "basic_salary"})


def create_employee_with_login(ctx: Ctx, data: EmployeeInput, role: Role, *, require_salary: bool) -> dict[str, Any]:
    """Creates User (temporary password, must change on first sign-in) + Employee (+ salary)."""
    _validate_input(ctx, data, require_salary=require_salary)
    temp_password = generate_temp_password()
    user = User(
        email=data.email.strip().lower(),
        full_name=data.full_name.strip(),
        password_hash=hash_password(temp_password, ctx.settings.bcrypt_rounds),
        role=role,
        is_active=True,
        must_change_password=True,
        created_at=ctx.now,
    )
    ctx.db.add(user)
    ctx.db.flush()
    employee = Employee(
        user_id=user.id,
        employee_code=(data.employee_code or "").strip() or _next_code(ctx),
        full_name=data.full_name.strip(),
        email=user.email,
        phone=data.phone,
        address=data.address,
        gender=(data.gender or None) and data.gender.strip().lower(),
        department_id=data.department_id,
        designation_id=data.designation_id,
        status=data.status,
        join_date=data.join_date,
        created_at=ctx.now,
    )
    ctx.db.add(employee)
    ctx.db.flush()
    if data.basic_salary:
        salary.add_structure(ctx, employee, data.basic_salary, data.join_date)
    audit.record(ctx, "create", "employee", employee.id, after=snapshot(employee))
    audit.record(ctx, "create", "user", user.id, after={"email": user.email, "role": user.role})
    result = serialize(employee)
    result["temporary_password"] = temp_password  # shown once to HR/Admin; user must change it
    return result


def serialize(e: Employee, *, include_login: bool = True) -> dict[str, Any]:
    data = {
        "id": e.id,
        "employee_code": e.employee_code,
        "full_name": e.full_name,
        "email": e.email,
        "phone": e.phone,
        "address": e.address,
        "gender": e.gender,
        "department": {"id": e.department_id, "name": e.department.name if e.department else None},
        "designation": {"id": e.designation_id, "name": e.designation.name} if e.designation else None,
        "status": e.status,
        "join_date": e.join_date.isoformat(),
        "exit_date": e.exit_date.isoformat() if e.exit_date else None,
    }
    if include_login and e.user is not None:
        data["login"] = {"user_id": e.user_id, "role": e.user.role, "is_active": e.user.is_active}
    return data


def get_employee_obj(ctx: Ctx, employee_id: int) -> Employee:
    e = ctx.db.get(Employee, employee_id)
    if e is None:
        raise NotFound("Employee not found")
    return e


# ------------------------------------------------------------------ registration & edits
def register_employee(ctx: Ctx, data: EmployeeInput) -> dict[str, Any]:
    """HR registers an employee with job details and salary; this creates the employee's login."""
    require_roles(ctx, Role.HR)
    return create_employee_with_login(ctx, data, Role.EMPLOYEE, require_salary=True)


EDITABLE_BY_HR = {"full_name", "email", "phone", "address", "gender", "department_id", "designation_id", "status", "join_date"}


def update_employee(ctx: Ctx, employee_id: int, changes: dict[str, Any]) -> dict[str, Any]:
    require_roles(ctx, Role.HR)
    e = get_employee_obj(ctx, employee_id)
    before = snapshot(e)
    unknown = set(changes) - EDITABLE_BY_HR
    if unknown:
        raise ValidationFailed(f"Fields not editable here: {', '.join(sorted(unknown))}")
    for key, value in changes.items():
        if value is None:
            continue
        if key == "email":
            email = str(value).strip().lower()
            if not EMAIL_RE.match(email):
                raise ValidationFailed("A valid email is required", details={"field": "email"})
            clash = ctx.db.scalar(select(User.id).where(User.email == email, User.id != e.user_id))
            if clash:
                raise Conflict(f"Email {email} is already registered")
            e.email = email
            e.user.email = email
        elif key == "department_id":
            if ctx.db.get(Department, int(value)) is None:
                raise ValidationFailed("Unknown department", details={"field": "department_id"})
            e.department_id = int(value)
        elif key == "designation_id":
            if ctx.db.get(Designation, int(value)) is None:
                raise ValidationFailed("Unknown designation", details={"field": "designation_id"})
            e.designation_id = int(value)
        elif key == "status":
            if value not in (EmploymentStatus.ACTIVE, EmploymentStatus.PROBATION):
                raise ValidationFailed("Use the exit action to mark an employee as exited")
            e.status = value
        elif key == "full_name":
            if not str(value).strip():
                raise ValidationFailed("full_name cannot be empty")
            e.full_name = str(value).strip()
            e.user.full_name = e.full_name
        else:
            setattr(e, key, value)
    ctx.db.flush()
    audit.record(ctx, "update", "employee", e.id, before=before, after=snapshot(e))
    return serialize(e)


def update_own_profile(ctx: Ctx, phone: str | None, address: str | None) -> dict[str, Any]:
    """Employees may edit only their own phone and address."""
    principal = ctx.actor
    if principal.employee_id is None:
        raise NotFound("You do not have an employee profile")
    e = get_employee_obj(ctx, principal.employee_id)
    before = snapshot(e)
    if phone is not None:
        e.phone = phone.strip() or None
    if address is not None:
        e.address = address.strip() or None
    audit.record(ctx, "update", "employee_profile", e.id, before=before, after=snapshot(e))
    return serialize(e)


def exit_employee(ctx: Ctx, employee_id: int, exit_date: date) -> dict[str, Any]:
    require_roles(ctx, Role.HR)
    e = get_employee_obj(ctx, employee_id)
    if e.status == EmploymentStatus.EXITED:
        raise Conflict("Employee has already exited")
    if exit_date < e.join_date:
        raise ValidationFailed("exit_date cannot be before join_date")
    before = snapshot(e)
    e.status = EmploymentStatus.EXITED
    e.exit_date = exit_date
    e.user.is_active = False
    e.user.token_version += 1
    audit.record(ctx, "update", "employee_exit", e.id, before=before, after=snapshot(e))
    return serialize(e)


# ------------------------------------------------------------------ directory & 360
def list_employees(
    ctx: Ctx,
    *,
    q: str | None = None,
    department_id: int | None = None,
    designation_id: int | None = None,
    status: str | None = None,
    page: int = 1,
    size: int = 20,
) -> dict[str, Any]:
    """Server-side search, filters and paging. Employees only ever see their own row."""
    principal = ctx.actor
    page, size = paginate(page, size)
    stmt = select(Employee).options(
        joinedload(Employee.department), joinedload(Employee.designation), joinedload(Employee.user)
    )
    if not principal.sees_all_employees:
        stmt = stmt.where(Employee.id == (principal.employee_id or -1))
    if q:
        like = f"%{q.strip().lower()}%"
        stmt = stmt.where(
            or_(
                func.lower(Employee.full_name).like(like),
                func.lower(Employee.email).like(like),
                func.lower(Employee.employee_code).like(like),
            )
        )
    if department_id:
        stmt = stmt.where(Employee.department_id == department_id)
    if designation_id:
        stmt = stmt.where(Employee.designation_id == designation_id)
    if status:
        stmt = stmt.where(Employee.status == status)
    total = ctx.db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = ctx.db.scalars(stmt.order_by(Employee.full_name).offset((page - 1) * size).limit(size)).unique().all()
    return {"items": [serialize(e) for e in rows], "total": total, "page": page, "size": size}


def get_employee(ctx: Ctx, employee_id: int) -> dict[str, Any]:
    assert_can_view_employee(ctx, employee_id)
    return serialize(get_employee_obj(ctx, employee_id))


def employee_360(ctx: Ctx, employee_id: int) -> dict[str, Any]:
    """Profile + attendance this month + leave balances + recent leave + payslips (+ salary for HR/Admin)."""
    from . import analytics, leave, payroll  # local import avoids import cycles

    assert_can_view_employee(ctx, employee_id)
    e = get_employee_obj(ctx, employee_id)
    data: dict[str, Any] = {"profile": serialize(e)}
    data["attendance_this_month"] = analytics.attendance_stats(ctx, period="this_month", employee_id=employee_id)
    data["leave_balances"] = leave.balances(ctx, employee_id=employee_id)
    data["recent_leave"] = leave.list_leaves(ctx, employee_id=employee_id, limit=5)["items"]
    data["payslips"] = payroll.list_payslips(ctx, employee_id=employee_id, limit=6)
    if ctx.actor.sees_all_employees:
        data["salary"] = salary.serialize(salary.current_structure(ctx, employee_id))
    return data


# ------------------------------------------------------------------ CSV import / export
CSV_COLUMNS = ["employee_code", "full_name", "email", "department", "designation", "join_date",
               "basic_salary", "gender", "phone", "address", "status"]
REQUIRED_CSV = ["full_name", "email", "department", "join_date", "basic_salary"]


MAX_UPLOAD_BYTES = 2_000_000


def _cell_text(value: Any) -> str:
    """Excel cell -> CSV text: dates become ISO dates, whole numbers lose the '.0'."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_upload(filename: str, raw: bytes) -> str:
    """Accept a CSV (UTF-8) or an Excel workbook (.xlsx, first sheet) and return CSV text for import_csv."""
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValidationFailed("File is too large (max 2 MB)")
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ValidationFailed("Old .xls files aren't supported — save the sheet as .xlsx or CSV")
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except (zipfile.BadZipFile, KeyError, ValueError, OSError, IndexError) as exc:
            raise ValidationFailed("Couldn't read the Excel file — is it a valid .xlsx?") from exc
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_text(c) for c in row]
            if any(cells):
                writer.writerow(cells)
        return buf.getvalue()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationFailed("CSV must be UTF-8 encoded (in Excel: Save As → CSV UTF-8)") from exc


def import_template(ctx: Ctx) -> str:
    """Header row for the employee import (HR). Valid department/designation names are listed in the UI."""
    require_roles(ctx, Role.HR)
    buf = io.StringIO()
    csv.writer(buf).writerow(CSV_COLUMNS)
    return buf.getvalue()


def import_csv(ctx: Ctx, content: str, *, dry_run: bool = True) -> dict[str, Any]:
    """Validates every row; the dry run returns row + field + reason for each problem.

    A real import is atomic: it runs only when the file has zero errors.
    """
    require_roles(ctx, Role.HR)
    reader = csv.DictReader(io.StringIO(content.lstrip("﻿")))
    if reader.fieldnames is None:
        raise ValidationFailed("CSV file is empty")
    headers = [h.strip().lower() for h in reader.fieldnames]
    missing_cols = [c for c in REQUIRED_CSV if c not in headers]
    if missing_cols:
        raise ValidationFailed("Missing required columns", details={"missing_columns": missing_cols})

    departments = {d.name.lower(): d for d in ctx.db.scalars(select(Department))}
    designations = {d.name.lower(): d for d in ctx.db.scalars(select(Designation))}
    existing_emails = set(ctx.db.scalars(select(User.email)))
    existing_codes = set(ctx.db.scalars(select(Employee.employee_code)))
    seen_emails: set[str] = set()
    seen_codes: set[str] = set()

    errors: list[dict[str, Any]] = []
    valid: list[tuple[int, EmployeeInput]] = []
    for row_no, raw in enumerate(reader, start=2):  # row 1 is the header
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw.items()}
        row_errors: list[dict[str, Any]] = []

        def err(field_name: str, reason: str, _row: int = row_no, _sink: list = row_errors) -> None:
            _sink.append({"row": _row, "field": field_name, "reason": reason})

        for col in REQUIRED_CSV:
            if not row.get(col):
                err(col, "required")
        email = row.get("email", "").lower()
        if email:
            if not EMAIL_RE.match(email):
                err("email", "invalid email")
            elif email in existing_emails:
                err("email", "already registered")
            elif email in seen_emails:
                err("email", "duplicate in file")
            seen_emails.add(email)
        code = row.get("employee_code", "")
        if code:
            if code in existing_codes:
                err("employee_code", "already exists")
            elif code in seen_codes:
                err("employee_code", "duplicate in file")
            seen_codes.add(code)
        dept = departments.get(row.get("department", "").lower()) if row.get("department") else None
        if row.get("department") and dept is None:
            err("department", f"unknown department '{row['department']}'")
        desig = None
        if row.get("designation"):
            desig = designations.get(row["designation"].lower())
            if desig is None:
                err("designation", f"unknown designation '{row['designation']}'")
        join: date | None = None
        if row.get("join_date"):
            try:
                join = date.fromisoformat(row["join_date"])
            except ValueError:
                err("join_date", "must be an ISO date (YYYY-MM-DD)")
        basic: float | None = None
        if row.get("basic_salary"):
            try:
                basic = float(row["basic_salary"].replace(",", ""))
                if basic <= 0:
                    err("basic_salary", "must be greater than 0")
            except ValueError:
                err("basic_salary", "must be a number")
        status = (row.get("status") or EmploymentStatus.ACTIVE).lower()
        if status not in (EmploymentStatus.ACTIVE, EmploymentStatus.PROBATION):
            err("status", "must be 'active' or 'probation'")

        if row_errors:
            errors.extend(row_errors)
        else:
            assert dept is not None and join is not None and basic is not None
            valid.append((row_no, EmployeeInput(
                full_name=row["full_name"], email=email, department_id=dept.id,
                designation_id=desig.id if desig else None, join_date=join, basic_salary=basic,
                employee_code=code or None, gender=row.get("gender") or None,
                phone=row.get("phone") or None, address=row.get("address") or None, status=status,
            )))

    report = {"dry_run": dry_run, "total_rows": len(valid) + len({e["row"] for e in errors}),
              "valid_rows": len(valid), "errors": errors}
    if dry_run:
        return report
    if errors:
        raise ValidationFailed("CSV has errors; nothing was imported", details=errors)
    created = [create_employee_with_login(ctx, data, Role.EMPLOYEE, require_salary=True) for _, data in valid]
    audit.record(ctx, "import", "employee_csv", None, after={"created": len(created)})
    report["created"] = [
        {"id": c["id"], "employee_code": c["employee_code"], "email": c["email"],
         "temporary_password": c["temporary_password"]}
        for c in created
    ]
    return report


def export_csv(ctx: Ctx) -> str:
    require_roles(ctx, Role.HR)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_COLUMNS + ["exit_date"])
    rows = ctx.db.scalars(
        select(Employee).options(joinedload(Employee.department), joinedload(Employee.designation))
        .order_by(Employee.employee_code)
    ).all()
    for e in rows:
        s = salary.current_structure(ctx, e.id)
        writer.writerow([
            e.employee_code, e.full_name, e.email, e.department.name,
            e.designation.name if e.designation else "", e.join_date.isoformat(),
            f"{s.basic:.2f}" if s else "", e.gender or "", e.phone or "", e.address or "", e.status,
            e.exit_date.isoformat() if e.exit_date else "",
        ])
    return buf.getvalue()
